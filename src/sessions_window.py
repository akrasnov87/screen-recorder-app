"""Окно со списком всех записей (сессий).

Изменения:
  • Сканирование папки sessions/ вынесено в отдельный QThread
    (SessionsScanThread) — UI остаётся отзывчивым на большом архиве.
  • Чтение текста через file_readers.
"""
from __future__ import annotations

import html
import json
import os
import shutil
from datetime import datetime
from typing import Any, Dict, List, Optional

from PySide6.QtCore import Qt, QThread, QUrl, Signal
from PySide6.QtGui import QAction, QDesktopServices, QKeySequence
from PySide6.QtWidgets import (
    QAbstractItemView, QDialog, QFileDialog, QFrame, QHBoxLayout, QHeaderView,
    QInputDialog, QLabel, QMenu, QMenuBar, QMessageBox, QProgressDialog,
    QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

from .file_readers import read_any_text, read_json_file
from .logger import get_logger
from .markdown_docx import markdown_to_docx
from .markdown_editor import MarkdownEditorDialog, MarkdownViewerDialog
from .markdown_to_bitrix import markdown_to_plain, markdown_to_plain_with_bb
from .metadata_dialog import MetadataDialog
from .task_queue import TaskQueue

log = get_logger(__name__)


# --- Статусы ---
STATUS_UPLOADED = "Сохранено"
STATUS_TRANSCRIBING = "Транскрибация"
STATUS_PROCESSED = "Обработан"
STATUS_ERROR = "Ошибка"
STATUS_UNKNOWN = "Неизвестно"


STATUS_COLORS = {
    STATUS_PROCESSED:    "#2E7D32",
    STATUS_UPLOADED:     "#37474F",
    STATUS_TRANSCRIBING: "#B8860B",
    STATUS_ERROR:        "#C62828",
    STATUS_UNKNOWN:      "#616161",
}


# Расширения видео/аудио, которые ищем в папке сессии.
_VIDEO_EXTS = (".mp4", ".mkv", ".mov", ".avi", ".webm", ".flv", ".wmv",
               ".mp3", ".wav", ".m4a", ".aac", ".opus", ".ogg")

# Расширения аудио, наличие которого считаем признаком обработки.
_AUDIO_EXTS = (".mp3", ".aac", ".wav", ".opus", ".ogg", ".m4a")


# ---------------------------------------------------------------------------
# Поток сканирования сессий
# ---------------------------------------------------------------------------
class SessionsScanThread(QThread):
    """Фоновое сканирование папки sessions/ и сбор метаданных записей."""

    finished_ok = Signal(list)
    failed = Signal(str)

    def __init__(self, sessions_root: str,
                 tasks_index: Dict[str, Dict[str, Any]],
                 parent=None) -> None:
        super().__init__(parent)
        self._sessions_root = sessions_root
        self._tasks_index = tasks_index

    def run(self) -> None:
        try:
            rows = self._collect(self._sessions_root, self._tasks_index)
            self.finished_ok.emit(rows)
        except Exception as exc:
            log.exception("Ошибка сканирования сессий: %s", exc)
            self.failed.emit(str(exc))

    # ------------------------------------------------------------------
    # Сбор данных
    # ------------------------------------------------------------------
    @staticmethod
    def _find_first_video(session_dir: str) -> str:
        for ext in _VIDEO_EXTS:
            p = os.path.join(session_dir, f"video{ext}")
            if os.path.exists(p):
                return p
        return ""

    @staticmethod
    def _find_first_audio(session_dir: str) -> str:
        for ext in _AUDIO_EXTS:
            p = os.path.join(session_dir, f"video{ext}")
            if os.path.exists(p):
                return p
        return ""

    @staticmethod
    def _find_prompt(session_dir: str) -> str:
        for fname in ("deepseek_prompt.docx", "deepseek_prompt.md",
                      "deepseek_prompt.txt"):
            p = os.path.join(session_dir, fname)
            if os.path.exists(p):
                return p
        return ""

    def _collect(
        self,
        sessions_root: str,
        tasks_index: Dict[str, Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        if not os.path.isdir(sessions_root):
            log.warning("Папка сессий не найдена: %s", sessions_root)
            return rows

        try:
            entries = sorted(os.listdir(sessions_root))
        except OSError as exc:
            log.error("Не удалось прочитать %s: %s", sessions_root, exc)
            return rows

        for name in entries:
            session_dir = os.path.join(sessions_root, name)
            if not os.path.isdir(session_dir):
                continue

            meta_path = os.path.join(session_dir, "session.json")
            meta = read_json_file(meta_path) or {}

            video_path = self._find_first_video(session_dir)
            has_video = bool(video_path)
            has_transcript = os.path.exists(
                os.path.join(session_dir, "video.txt")
            )
            audio_path = self._find_first_audio(session_dir)
            has_audio = bool(audio_path)

            task = (tasks_index.get(os.path.abspath(video_path))
                    if video_path else None)
            task_id = (task or {}).get("task_id", "")

            if task is not None:
                st = task.get("status", "")
                if st == "completed":
                    status = STATUS_PROCESSED
                elif st == "error":
                    status = STATUS_ERROR
                elif st in ("pending", "converting", "transcribing",
                            "summarizing"):
                    status = STATUS_TRANSCRIBING
                else:
                    status = STATUS_UNKNOWN
            else:
                if has_transcript or has_audio:
                    status = STATUS_PROCESSED
                elif has_video:
                    status = STATUS_UPLOADED
                else:
                    status = STATUS_UNKNOWN

            date_str = meta.get("date") or ""
            time_str = meta.get("time") or ""
            try:
                parts = name.split("_")
                if len(parts) >= 2:
                    date_str = date_str or parts[0]
                    time_str = time_str or parts[1].replace("-", ":")
            except Exception:
                pass
            if not date_str:
                try:
                    ts = os.path.getmtime(session_dir)
                    date_str = datetime.fromtimestamp(ts).strftime(
                        "%Y-%m-%d"
                    )
                    time_str = datetime.fromtimestamp(ts).strftime(
                        "%H:%M:%S"
                    )
                except Exception:
                    pass

            prompt_path = self._find_prompt(session_dir)

            attachments = list(meta.get("attachments", []) or [])
            attachments_dir = os.path.join(session_dir, "attachments")
            if not attachments and os.path.isdir(attachments_dir):
                try:
                    attachments = [
                        os.path.join(attachments_dir, f)
                        for f in sorted(os.listdir(attachments_dir))
                        if os.path.isfile(
                            os.path.join(attachments_dir, f)
                        )
                    ]
                except OSError:
                    attachments = []

            manual_protocol_path = meta.get("manual_protocol_path") or ""
            if manual_protocol_path and not os.path.exists(
                manual_protocol_path
            ):
                manual_protocol_path = ""

            summary_bb = str(meta.get("summary_bb") or "")

            rows.append({
                "dir": session_dir,
                "name": meta.get("name") or name,
                "project": meta.get("project") or "",
                "status": status,
                "source": meta.get("source") or "record",
                "is_scrum": bool(meta.get("is_scrum", False)),
                "datetime": f"{date_str} {time_str}".strip(),
                "video_path": video_path,
                "has_video": has_video,
                "task_id": task_id,
                "prompt_path": prompt_path,
                "attachments": attachments,
                "manual_protocol_path": manual_protocol_path,
                "summary_bb": summary_bb,
            })

        rows.sort(key=lambda r: r["datetime"], reverse=True)
        return rows


class SessionsWindow(QDialog):
    """Список всех записей в папке sessions/."""

    import_requested = Signal()

    def __init__(
        self,
        sessions_root: str,
        task_queue: TaskQueue,
        processor,
        config_manager=None,
        parent: Optional[QWidget] = None,
    ) -> None:
        super().__init__(parent)
        self.sessions_root = sessions_root
        self.task_queue = task_queue
        self.processor = processor
        self.config_manager = config_manager
        self.setWindowTitle("Записи")
        self.setMinimumSize(1400, 780)
        self.setModal(False)
        self._rows: List[Dict[str, Any]] = []
        self._thread: Optional[SessionsScanThread] = None

        self._build_ui()
        self._build_menu_bar()
        self.refresh()

    # ------------------------------------------------------------------
    # UI
    # ------------------------------------------------------------------
    def _build_ui(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(8, 4, 8, 8)
        root.setSpacing(6)

        header = QHBoxLayout()
        self.summary_label = QLabel("")
        header.addWidget(self.summary_label)
        header.addSpacing(20)
        self.selection_label = QLabel("")
        self.selection_label.setStyleSheet("QLabel { color: #666; }")
        header.addWidget(self.selection_label)
        header.addStretch()

        self.refresh_indicator = QLabel("")
        self.refresh_indicator.setStyleSheet(
            "QLabel { color: #4a90d9; font-style: italic; }"
        )
        header.addWidget(self.refresh_indicator)
        root.addLayout(header)

        self.table = QTableWidget(0, 10)
        self.table.setHorizontalHeaderLabels([
            "Дата и время", "Название", "Проект", "Статус",
            "Источник", "Скрам", "Вложения", "Summary",
            "Task ID", "Папка",
        ])
        self.table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )
        self.table.setEditTriggers(
            QAbstractItemView.EditTrigger.NoEditTriggers
        )
        hv = self.table.horizontalHeader()
        hv.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        hv.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(6, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(7, QHeaderView.ResizeMode.Stretch)
        hv.setSectionResizeMode(8, QHeaderView.ResizeMode.ResizeToContents)
        hv.setSectionResizeMode(9, QHeaderView.ResizeMode.Stretch)
        self.table.itemSelectionChanged.connect(
            self._on_selection_changed
        )
        self.table.setContextMenuPolicy(
            Qt.ContextMenuPolicy.CustomContextMenu
        )
        self.table.customContextMenuRequested.connect(
            self._show_context_menu
        )
        root.addWidget(self.table, 1)

        root.addWidget(self._build_legend())

        bottom = QHBoxLayout()
        bottom.addStretch()

        self.refresh_btn = QPushButton("Обновить")
        self.refresh_btn.clicked.connect(self.refresh)
        bottom.addWidget(self.refresh_btn)

        self.close_btn = QPushButton("Закрыть")
        self.close_btn.clicked.connect(self.close)
        bottom.addWidget(self.close_btn)

        root.addLayout(bottom)

    def _build_menu_bar(self) -> None:
        bar = QMenuBar(self)
        _MENU_QSS = (
            "QMenuBar {"
            "  background-color: palette(window);"
            "  color: palette(window-text);"
            "}"
            "QMenuBar::item {"
            "  background: transparent;"
            "  color: palette(window-text);"
            "  padding: 4px 10px;"
            "}"
            "QMenuBar::item:selected {"
            "  background-color: #2D7FF9;"
            "  color: #FFFFFF;"
            "}"
            "QMenuBar::item:pressed {"
            "  background-color: #1E5FBF;"
            "  color: #FFFFFF;"
            "}"
            "QMenu {"
            "  background-color: palette(window);"
            "  color: palette(window-text);"
            "  border: 1px solid palette(mid);"
            "}"
            "QMenu::item {"
            "  background: transparent;"
            "  color: palette(window-text);"
            "  padding: 5px 24px 5px 24px;"
            "}"
            "QMenu::item:selected {"
            "  background-color: #2D7FF9;"
            "  color: #FFFFFF;"
            "}"
            "QMenu::item:disabled {"
            "  color: palette(mid);"
            "}"
            "QMenu::separator {"
            "  height: 1px;"
            "  background-color: palette(mid);"
            "  margin: 4px 8px;"
            "}"
        )
        bar.setStyleSheet(_MENU_QSS)
        self.setStyleSheet(self.styleSheet() + _MENU_QSS)

        layout: QVBoxLayout = self.layout()
        layout.insertWidget(0, bar)

        # ---------------- Файл ----------------
        m_file = bar.addMenu("Файл")

        act_import = QAction("Импорт материалов…", self)
        act_import.setShortcut(QKeySequence("Ctrl+I"))
        act_import.setToolTip(
            "Импортировать готовое видео, стенограмму или протокол"
        )
        act_import.triggered.connect(self._on_import_requested)
        m_file.addAction(act_import)

        m_file.addSeparator()

        act_open_folder = QAction("Открыть папку записи", self)
        act_open_folder.setShortcut(QKeySequence("Ctrl+Shift+E"))
        act_open_folder.triggered.connect(self._open_folder)
        m_file.addAction(act_open_folder)

        act_open_video = QAction("Открыть видео", self)
        act_open_video.setShortcut(QKeySequence("Ctrl+Shift+V"))
        act_open_video.triggered.connect(self._open_video)
        m_file.addAction(act_open_video)

        m_file.addSeparator()

        act_delete = QAction("Удалить запись…", self)
        act_delete.setShortcut(QKeySequence("Ctrl+Delete"))
        act_delete.triggered.connect(self._delete_session)
        m_file.addAction(act_delete)

        m_file.addSeparator()

        act_refresh = QAction("Обновить список", self)
        act_refresh.setShortcut(QKeySequence("F5"))
        act_refresh.triggered.connect(self.refresh)
        m_file.addAction(act_refresh)

        act_close = QAction("Закрыть окно", self)
        act_close.setShortcut(QKeySequence("Ctrl+W"))
        act_close.triggered.connect(self.close)
        m_file.addAction(act_close)

        # ---------------- Протокол ----------------
        m_protocol = bar.addMenu("Протокол")

        act_edit_protocol_md = QAction(
            "Создать/редактировать протокол (Markdown)…", self
        )
        act_edit_protocol_md.setShortcut(QKeySequence("Ctrl+M"))
        act_edit_protocol_md.setToolTip(
            "Открыть встроенный Markdown-редактор. Слева — исходный "
            "текст, справа — готовый документ.\n"
            "Файл сохраняется как manual_protocol.md и прикрепляется "
            "к записи."
        )
        act_edit_protocol_md.triggered.connect(self._edit_manual_protocol_md)
        m_protocol.addAction(act_edit_protocol_md)

        act_export_protocol_docx = QAction(
            "Экспорт протокола в DOCX…", self
        )
        act_export_protocol_docx.setShortcut(QKeySequence("Ctrl+Shift+M"))
        act_export_protocol_docx.setToolTip(
            "Сконвертировать прикреплённый протокол (Markdown или "
            "текст) в документ .docx.\n\n"
            "По умолчанию результат сохраняется как "
            "<session_dir>/manual_protocol.docx и прикрепляется к "
            "записи — после этого «Отправить в чат…» (Bitrix24) "
            "сможет отправить именно DOCX."
        )
        act_export_protocol_docx.triggered.connect(
            self._export_protocol_docx
        )
        m_protocol.addAction(act_export_protocol_docx)

        m_protocol.addSeparator()

        act_attach_protocol = QAction("Прикрепить файл протокола…", self)
        act_attach_protocol.setToolTip(
            "Загрузить вручную подготовленный протокол "
            "(.docx/.txt/.md/.pdf) и прикрепить его к записи"
        )
        act_attach_protocol.triggered.connect(self._attach_manual_protocol)
        m_protocol.addAction(act_attach_protocol)

        act_open_protocol = QAction("Открыть прикреплённый протокол", self)
        act_open_protocol.triggered.connect(self._open_manual_protocol)
        m_protocol.addAction(act_open_protocol)

        # ---------------- Summary ----------------
        m_summary = bar.addMenu("Summary")

        act_edit_summary = QAction("Изменить summary (Markdown)…", self)
        act_edit_summary.setShortcut(QKeySequence("Ctrl+P"))
        act_edit_summary.setToolTip(
            "Открыть Markdown-редактор для краткого описания записи.\n"
            "При отправке в Bitrix24 Markdown конвертируется в BB-код."
        )
        act_edit_summary.triggered.connect(self._edit_summary_bb)
        m_summary.addAction(act_edit_summary)

        act_view_summary = QAction("Просмотр summary", self)
        act_view_summary.setShortcut(QKeySequence("Ctrl+Shift+P"))
        act_view_summary.setToolTip(
            "Открыть краткое описание в режиме только для чтения "
            "(с отрендеренным Markdown)"
        )
        act_view_summary.triggered.connect(self._view_summary)
        m_summary.addAction(act_view_summary)

        act_export_summary = QAction("Экспорт summary…", self)
        act_export_summary.setToolTip(
            "Сохранить краткое описание в .docx / .html / .md / .txt"
        )
        act_export_summary.triggered.connect(self._export_summary)
        m_summary.addAction(act_export_summary)

        # ---------------- DeepSeek ----------------
        m_deepseek = bar.addMenu("DeepSeek")

        act_open_prompt = QAction("Открыть промпт DeepSeek", self)
        act_open_prompt.setShortcut(QKeySequence("Ctrl+D"))
        act_open_prompt.triggered.connect(self._open_deepseek_prompt)
        m_deepseek.addAction(act_open_prompt)

        act_export_prompt = QAction("Экспорт промпта…", self)
        act_export_prompt.setShortcut(QKeySequence("Ctrl+Shift+D"))
        act_export_prompt.triggered.connect(self._export_prompt)
        m_deepseek.addAction(act_export_prompt)

        m_deepseek.addSeparator()

        act_save_downloads = QAction("Сохранить промпт в Загрузки", self)
        act_save_downloads.setShortcut(QKeySequence("Ctrl+Alt+D"))
        act_save_downloads.setToolTip(
            "Скопировать файлы промпта (deepseek_prompt.*) в папку "
            "«Загрузки» без запроса пути"
        )
        act_save_downloads.triggered.connect(self._save_prompt_to_downloads)
        m_deepseek.addAction(act_save_downloads)

        # ---------------- Bitrix24 ----------------
        m_bitrix = bar.addMenu("Bitrix24")

        act_send = QAction("Отправить в чат…", self)
        act_send.setShortcut(QKeySequence("Ctrl+B"))
        act_send.setToolTip(
            "Отправить протокол и/или summary записи в чат Bitrix24"
        )
        act_send.triggered.connect(self._send_to_bitrix)
        m_bitrix.addAction(act_send)

        act_send_protocol = QAction("Отправить протокол…", self)
        act_send_protocol.triggered.connect(
            lambda: self._send_to_bitrix(default="protocol")
        )
        m_bitrix.addAction(act_send_protocol)

        act_send_summary = QAction("Отправить summary…", self)
        act_send_summary.triggered.connect(
            lambda: self._send_to_bitrix(default="summary")
        )
        m_bitrix.addAction(act_send_summary)

        # ---------------- Вложения ----------------
        m_attach = bar.addMenu("Вложения")

        act_open_attachments = QAction("Открыть папку вложений", self)
        act_open_attachments.triggered.connect(self._open_attachments)
        m_attach.addAction(act_open_attachments)

        act_add_attachment = QAction("Добавить вложение…", self)
        act_add_attachment.triggered.connect(self._add_attachment_to_session)
        m_attach.addAction(act_add_attachment)

        # ---------------- Очередь и статус ----------------
        m_queue = bar.addMenu("Очередь")

        act_edit_meta = QAction(
            "Редактировать метаданные и перезапустить…", self
        )
        act_edit_meta.setShortcut(QKeySequence("Ctrl+E"))
        act_edit_meta.setToolTip(
            "Открыть карточку записи с текущими параметрами.\n"
            "Изменённые настройки (скрам, промпт DeepSeek, вложения, "
            "контекст в промпте) сохранятся в session.json, и запись "
            "будет поставлена в очередь на повторную обработку.\n\n"
            "Удобно, если при первичной обработке забыли включить "
            "«Сформировать файл промпта для DeepSeek»."
        )
        act_edit_meta.triggered.connect(self._edit_metadata_and_restart)
        m_queue.addAction(act_edit_meta)

        m_queue.addSeparator()

        act_restart = QAction("Перезапустить обработку", self)
        act_restart.setShortcut(QKeySequence("Ctrl+R"))
        act_restart.setToolTip(
            "Перезапустить обработку с текущими метаданными "
            "(без открытия карточки)"
        )
        act_restart.triggered.connect(self._restart_processing)
        m_queue.addAction(act_restart)

        act_enqueue = QAction("Поставить в очередь", self)
        act_enqueue.triggered.connect(self._enqueue_current)
        m_queue.addAction(act_enqueue)

        m_queue.addSeparator()

        act_status_uploaded = QAction("Пометить: «Сохранено»", self)
        act_status_uploaded.triggered.connect(
            lambda: self._change_status(STATUS_UPLOADED)
        )
        m_queue.addAction(act_status_uploaded)

        act_status_processed = QAction("Пометить: «Обработан»", self)
        act_status_processed.triggered.connect(
            lambda: self._change_status(STATUS_PROCESSED)
        )
        m_queue.addAction(act_status_processed)

        act_status_error = QAction("Пометить: «Ошибка»", self)
        act_status_error.triggered.connect(
            lambda: self._change_status(STATUS_ERROR)
        )
        m_queue.addAction(act_status_error)

        # ---------------- Утилиты ----------------
        m_utils = bar.addMenu("Утилиты")

        act_stats = QAction("Показать размер видеофайлов…", self)
        act_stats.setToolTip(
            "Посчитать, сколько места занимают video.mp4 во всех сессиях"
        )
        act_stats.triggered.connect(self._show_video_stats)
        m_utils.addAction(act_stats)

        m_utils.addSeparator()

        act_del_video = QAction(
            "Удалить видеофайлы (оставить только аудио)…", self
        )
        act_del_video.setToolTip(
            "Удалить video.mp4 из всех сессий, где уже есть аудиофайл.\n"
            "Помогает освободить дисковое пространство.\n"
            "Требует подтверждения."
        )
        act_del_video.triggered.connect(self._delete_video_files)
        m_utils.addAction(act_del_video)

        # ---------------- Справка ----------------
        m_help = bar.addMenu("Справка")

        act_shortcuts = QAction("Горячие клавиши…", self)
        act_shortcuts.triggered.connect(self._show_shortcuts)
        m_help.addAction(act_shortcuts)

    # ------------------------------------------------------------------
    # Сигнал импорта
    # ------------------------------------------------------------------
    def _on_import_requested(self) -> None:
        log.debug("Запрошен импорт из окна «Записи»")
        self.import_requested.emit()

    # ------------------------------------------------------------------
    # Легенда
    # ------------------------------------------------------------------
    def _build_legend(self) -> QWidget:
        box = QFrame()
        box.setFrameShape(QFrame.Shape.StyledPanel)
        box.setStyleSheet(
            "QFrame { background-color: palette(window); "
            "border: 1px solid palette(mid); border-radius: 6px; }"
        )
        layout = QHBoxLayout(box)
        layout.setContentsMargins(10, 4, 10, 4)
        layout.setSpacing(14)

        layout.addWidget(QLabel("<b>Статусы:</b>"))
        items = [
            (STATUS_UPLOADED, "видео сохранено, обработка не выполнялась"),
            (STATUS_TRANSCRIBING,
             "задача в очереди: конвертация/транскрибация"),
            (STATUS_PROCESSED,
             "все шаги завершены — есть audio и/или transcript"),
            (STATUS_ERROR,
             "задача завершилась с ошибкой; доступен ручной перезапуск"),
            (STATUS_UNKNOWN, "не удалось определить состояние"),
        ]
        for status, hint in items:
            layout.addWidget(self._make_legend_item(status, hint))
        layout.addStretch()
        return box

    @staticmethod
    def _make_legend_item(status: str, hint: str) -> QWidget:
        w = QWidget()
        w.setToolTip(hint)
        row = QHBoxLayout(w)
        row.setContentsMargins(0, 0, 0, 0)
        row.setSpacing(6)
        swatch = QLabel()
        swatch.setFixedSize(12, 12)
        swatch.setStyleSheet(
            f"background-color: {STATUS_COLORS.get(status, '#616161')}; "
            f"border-radius: 2px;"
        )
        row.addWidget(swatch)
        text = QLabel(status)
        row.addWidget(text)
        hint_lbl = QLabel("?")
        hint_lbl.setToolTip(hint)
        hint_lbl.setStyleSheet(
            "color: gray; font-weight: bold; padding-left: 2px;"
        )
        row.addWidget(hint_lbl)
        return w

    # ------------------------------------------------------------------
    # Выбор строки
    # ------------------------------------------------------------------
    def _on_selection_changed(self) -> None:
        r = self._selected_row()
        if not r:
            self.selection_label.setText("")
            return
        parts = [f"<b>{r['name']}</b>"]
        if r.get("project"):
            parts.append(f"проект: {r['project']}")
        parts.append(f"статус: {r['status']}")
        if r.get("manual_protocol_path"):
            parts.append("протокол: прикреплён")
        if (r.get("summary_bb") or "").strip():
            parts.append("summary: есть")
        self.selection_label.setText(" | ".join(parts))

    # ------------------------------------------------------------------
    # Обновление
    # ------------------------------------------------------------------
    def _build_tasks_index(self) -> Dict[str, Dict[str, Any]]:
        tasks_index: Dict[str, Dict[str, Any]] = {}
        try:
            for t in self.task_queue.get_queue():
                vp = t.get("video_path", "")
                if vp:
                    tasks_index[os.path.abspath(vp)] = t
        except Exception as exc:
            log.warning("Не удалось получить очередь: %s", exc)
        return tasks_index

    def refresh(self) -> None:
        if self._thread is not None and self._thread.isRunning():
            log.debug("Сканирование уже идёт, пропускаем refresh")
            return

        self.refresh_indicator.setText("Сканирование…")
        self.refresh_btn.setEnabled(False)

        tasks_index = self._build_tasks_index()
        self._thread = SessionsScanThread(
            self.sessions_root, tasks_index, parent=self
        )
        self._thread.finished_ok.connect(self._on_scan_finished)
        self._thread.failed.connect(self._on_scan_failed)
        self._thread.finished.connect(self._on_scan_thread_done)
        self._thread.start()

    def _on_scan_finished(self, rows: list) -> None:
        self._rows = list(rows or [])
        self._render_rows()
        self._update_summary()

    def _on_scan_failed(self, error: str) -> None:
        log.error("Сканирование сессий провалено: %s", error)
        QMessageBox.warning(self, "Записи",
                            f"Ошибка сканирования:\n{error}")

    def _on_scan_thread_done(self) -> None:
        self.refresh_indicator.setText("")
        self.refresh_btn.setEnabled(True)
        self._thread = None

    def _update_summary(self) -> None:
        rows = self._rows
        total = len(rows)
        processed = sum(1 for r in rows
                        if r["status"] == STATUS_PROCESSED)
        uploaded = sum(1 for r in rows
                       if r["status"] == STATUS_UPLOADED)
        in_progress = sum(1 for r in rows
                          if r["status"] == STATUS_TRANSCRIBING)
        errors = sum(1 for r in rows if r["status"] == STATUS_ERROR)

        self.summary_label.setText(
            f"Всего: {total} | Обработан: {processed} | "
            f"Сохранено: {uploaded} | В обработке: {in_progress} | "
            f"Ошибок: {errors}"
        )
        log.debug("Список записей обновлён: %d сессий", total)
        self._on_selection_changed()

    def _render_rows(self) -> None:
        self.table.setRowCount(0)
        for r in self._rows:
            row = self.table.rowCount()
            self.table.insertRow(row)

            self.table.setItem(row, 0, QTableWidgetItem(r["datetime"]))
            self.table.setItem(row, 1, QTableWidgetItem(r["name"]))
            self.table.setItem(row, 2, QTableWidgetItem(r["project"]))

            status_item = QTableWidgetItem(r["status"])
            if r["status"] == STATUS_ERROR:
                status_item.setForeground(Qt.GlobalColor.red)
            elif r["status"] == STATUS_PROCESSED:
                status_item.setForeground(Qt.GlobalColor.darkGreen)
            elif r["status"] == STATUS_TRANSCRIBING:
                status_item.setForeground(Qt.GlobalColor.darkYellow)
            self.table.setItem(row, 3, status_item)

            src_map = {
                "record": "Запись",
                "upload": "Загрузка",
                "import": "Импорт",
            }
            src = src_map.get(r["source"], r["source"] or "—")
            self.table.setItem(row, 4, QTableWidgetItem(src))
            self.table.setItem(row, 5,
                               QTableWidgetItem("да" if r["is_scrum"]
                                                else "—"))

            att_count = len(r.get("attachments", []) or [])
            self.table.setItem(
                row, 6,
                QTableWidgetItem(str(att_count) if att_count else "—")
            )

            summary_bb = (r.get("summary_bb") or "").strip()
            if summary_bb:
                short = summary_bb.replace("\n", " ")
                if len(short) > 60:
                    short = short[:57] + "…"
                summary_item = QTableWidgetItem(short)
                summary_item.setToolTip(summary_bb[:1000])
            else:
                summary_item = QTableWidgetItem("—")
            self.table.setItem(row, 7, summary_item)

            self.table.setItem(row, 8,
                               QTableWidgetItem(r["task_id"] or "—"))
            self.table.setItem(row, 9, QTableWidgetItem(r["dir"]))

    # ------------------------------------------------------------------
    # Bitrix24
    # ------------------------------------------------------------------
    def _send_to_bitrix(self, default: str = "") -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        if self.config_manager is None:
            QMessageBox.warning(
                self, "Bitrix24",
                "Нет доступа к настройкам — ConfigManager не передан.",
            )
            return

        bitrix_cfg = self.config_manager.get_bitrix_settings()
        if not bitrix_cfg.get("enabled"):
            QMessageBox.information(
                self, "Bitrix24",
                "Интеграция с Bitrix24 отключена.\n\n"
                "Включите её в Настройки → Bitrix24.",
            )
            return

        project = r.get("project") or ""
        chat_id = self.config_manager.get_project_chat_id(project)
        if not chat_id:
            log.info(
                "Для проекта «%s» не задан ID чата — "
                "пользователь выберет получателя вручную",
                project or "—",
            )

        session_info = {
            "name": r.get("name") or "",
            "project": project,
            "date": r.get("datetime") or "",
            "protocol_path": r.get("manual_protocol_path") or "",
            "protocol_label": (
                os.path.basename(r["manual_protocol_path"])
                if r.get("manual_protocol_path") else ""
            ),
            "summary_bb": r.get("summary_bb") or "",
            "comment": "",
            "session_dir": r.get("dir") or "",
        }

        try:
            meta = read_json_file(
                os.path.join(r["dir"], "session.json")
            ) or {}
            session_info["comment"] = meta.get("comment") or ""
        except Exception:
            pass

        from .send_to_bitrix_dialog import SendToBitrixDialog

        projects = self.config_manager.get_projects()
        employees = self.config_manager.get_employees()

        dlg = SendToBitrixDialog(
            session_info=session_info,
            chat_id=chat_id,
            bitrix_cfg=bitrix_cfg,
            projects=projects,
            employees=employees,
            parent=self,
        )
        dlg.exec()

    # ------------------------------------------------------------------
    # Протокол
    # ------------------------------------------------------------------
    def _edit_manual_protocol_md(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        md_path = os.path.join(r["dir"], "manual_protocol.md")
        initial_text = ""

        if os.path.exists(md_path):
            initial_text = read_any_text(md_path)
        else:
            existing = r.get("manual_protocol_path") or ""
            if existing and os.path.exists(existing):
                initial_text = self._read_protocol_as_text(existing)

        default_docx = os.path.join(r["dir"], "manual_protocol.docx")

        dlg = MarkdownEditorDialog(
            text=initial_text,
            title=f"Протокол (Markdown) — {r['name']}",
            parent=self,
            default_docx_path=default_docx,
        )
        if dlg.exec() != QDialog.DialogCode.Accepted:
            log.debug("Markdown-редактор протокола закрыт без сохранения")
            return

        new_text = dlg.result_text()

        try:
            with open(md_path, "w", encoding="utf-8") as f:
                f.write(new_text)
            log.info("Markdown-протокол сохранён: %s (%d символов)",
                     md_path, len(new_text))
        except Exception as exc:
            log.exception("Не удалось записать %s: %s", md_path, exc)
            QMessageBox.critical(
                self, "Протокол",
                f"Не удалось сохранить протокол:\n{exc}",
            )
            return

        session_json = os.path.join(r["dir"], "session.json")
        meta = read_json_file(session_json) or {}
        meta["manual_protocol_path"] = md_path
        if not self._write_json(session_json, meta):
            QMessageBox.critical(
                self, "Протокол",
                "Файл протокола записан, но не удалось обновить "
                "session.json.",
            )
            return

        self.refresh()
        QMessageBox.information(
            self, "Протокол",
            f"Протокол сохранён и прикреплён к записи:\n{md_path}\n\n"
            "Чтобы получить .docx для отправки в чат — "
            "используйте «Протокол → Экспорт протокола в DOCX…».",
        )

    def _export_protocol_docx(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        src_path = r.get("manual_protocol_path") or ""
        md_path = os.path.join(r["dir"], "manual_protocol.md")

        if not src_path or not os.path.exists(src_path):
            if os.path.exists(md_path):
                src_path = md_path
            else:
                QMessageBox.information(
                    self, "Экспорт в DOCX",
                    "К записи не прикреплён протокол.\n\n"
                    "Используйте «Протокол → Создать/редактировать "
                    "протокол (Markdown)…» или «Прикрепить файл "
                    "протокола…».",
                )
                return

        ext = os.path.splitext(src_path)[1].lower()

        if ext == ".docx":
            QMessageBox.information(
                self, "Экспорт в DOCX",
                f"Протокол уже в формате DOCX:\n{src_path}",
            )
            return

        md_text = self._read_protocol_as_text(src_path)
        if not md_text.strip():
            QMessageBox.warning(
                self, "Экспорт в DOCX",
                "Не удалось извлечь текст из исходного файла.\n"
                "Поддерживаются .md, .txt, .docx, .pdf.",
            )
            return

        default_path = os.path.join(r["dir"], "manual_protocol.docx")
        target, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить протокол как DOCX",
            default_path,
            "Документы Word (*.docx);;Все файлы (*)",
        )
        if not target:
            return

        if not target.lower().endswith(".docx"):
            target += ".docx"

        title = r.get("name") or ""

        try:
            markdown_to_docx(md_text, target, title=title)
        except Exception as exc:
            log.exception("Ошибка конвертации в DOCX: %s", exc)
            QMessageBox.critical(
                self, "Экспорт в DOCX",
                f"Не удалось сохранить файл:\n{exc}",
            )
            return

        attach = False
        if os.path.abspath(target) == os.path.abspath(default_path):
            session_json = os.path.join(r["dir"], "session.json")
            meta = read_json_file(session_json) or {}
            meta["manual_protocol_path"] = target
            if self._write_json(session_json, meta):
                attach = True

        if attach:
            log.info("Протокол DOCX сохранён и прикреплён к записи: %s",
                     target)
            QMessageBox.information(
                self, "Экспорт в DOCX",
                f"Документ сохранён и прикреплён к записи:\n{target}\n\n"
                "Теперь его можно отправить в Bitrix24 через "
                "«Bitrix24 → Отправить в чат…».",
            )
            self.refresh()
        else:
            log.info("Протокол DOCX сохранён отдельно: %s", target)
            QMessageBox.information(
                self, "Экспорт в DOCX",
                f"Документ сохранён:\n{target}",
            )

    @staticmethod
    def _write_json(path: str, data: Dict[str, Any]) -> bool:
        try:
            tmp = path + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            os.replace(tmp, path)
            return True
        except Exception as exc:
            log.exception("Не удалось записать %s: %s", path, exc)
            return False

    @staticmethod
    def _read_protocol_as_text(path: str) -> str:
        return read_any_text(path)

    def _attach_manual_protocol(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите файл протокола",
            os.path.expanduser("~"),
            "Документы (*.docx *.txt *.md *.pdf);;Все файлы (*)",
        )
        if not file_path:
            return

        if not os.path.isfile(file_path):
            QMessageBox.warning(self, "Протокол",
                                f"Файл не найден:\n{file_path}")
            return

        ext = os.path.splitext(file_path)[1].lower() or ".bin"
        target = os.path.join(r["dir"], f"manual_protocol{ext}")

        try:
            shutil.copy2(file_path, target)
            log.info("Ручной протокол скопирован: %s → %s",
                     file_path, target)
        except Exception as exc:
            log.exception("Ошибка копирования протокола: %s", exc)
            QMessageBox.critical(self, "Протокол",
                                 f"Не удалось скопировать файл:\n{exc}")
            return

        session_json = os.path.join(r["dir"], "session.json")
        meta = read_json_file(session_json) or {}
        meta["manual_protocol_path"] = target
        if not self._write_json(session_json, meta):
            QMessageBox.critical(self, "Протокол",
                                 "Не удалось обновить session.json")
            return

        QMessageBox.information(
            self, "Протокол",
            f"Протокол прикреплён к записи:\n{target}",
        )
        self.refresh()

    def _open_manual_protocol(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        path = r.get("manual_protocol_path") or ""
        if not path or not os.path.exists(path):
            QMessageBox.information(
                self, "Протокол",
                "К этой записи не прикреплён ручной протокол.\n\n"
                "Меню «Протокол» → «Создать/редактировать протокол "
                "(Markdown)…» или «Прикрепить файл протокола…».",
            )
            return
        log.info("Открытие ручного протокола: %s", path)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    # ------------------------------------------------------------------
    # Summary (Markdown)
    # ------------------------------------------------------------------
    def _edit_summary_bb(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        dlg = MarkdownEditorDialog(
            text=r.get("summary_bb") or "",
            title=f"Краткое описание — {r['name']}",
            parent=self,
        )
        if dlg.exec() != QDialog.DialogCode.Accepted:
            log.debug("Markdown-редактор summary закрыт без сохранения")
            return

        new_text = dlg.result_text()
        session_json = os.path.join(r["dir"], "session.json")
        meta = read_json_file(session_json) or {}
        meta["summary_bb"] = new_text
        if not self._write_json(session_json, meta):
            QMessageBox.critical(
                self, "Summary",
                "Не удалось сохранить summary в session.json",
            )
            return

        try:
            summary_md_path = os.path.join(r["dir"], "summary.md")
            with open(summary_md_path, "w", encoding="utf-8") as f:
                f.write(new_text)
            log.info("Summary-Markdown сохранён: %s (%d символов)",
                     summary_md_path, len(new_text))
        except Exception as exc:
            log.warning("Не удалось сохранить summary.md: %s", exc)

        log.info("Summary обновлён для %s (%d символов)",
                 r["dir"], len(new_text))
        self.refresh()

    def _view_summary(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        text = r.get("summary_bb") or ""
        if not text.strip():
            QMessageBox.information(
                self, "Summary",
                "У этой записи ещё нет краткого описания.\n\n"
                "Меню «Summary» → «Изменить summary (Markdown)…»",
            )
            return
        dlg = MarkdownViewerDialog(
            text=text,
            title=f"Просмотр summary — {r['name']}",
            parent=self,
        )
        dlg.exec()

    def _export_summary(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        text_md = r.get("summary_bb") or ""
        if not text_md.strip():
            QMessageBox.information(
                self, "Экспорт summary",
                "У этой записи нет краткого описания.",
            )
            return

        formats = ["docx", "html", "md", "txt"]
        fmt, ok = QInputDialog.getItem(
            self, "Экспорт summary", "Формат файла:",
            formats, 0, False,
        )
        if not ok or not fmt:
            return

        default_path = os.path.join(
            r["dir"], f"summary_{self._safe_name(r['name'])}.{fmt}"
        )
        target_path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить summary как",
            default_path,
            "Word (*.docx);;HTML (*.html);;Markdown (*.md);;"
            "Text (*.txt);;All files (*)",
        )
        if not target_path:
            return

        try:
            if fmt == "docx":
                markdown_to_docx(
                    text_md, target_path, title=r.get("name") or ""
                )
            elif fmt == "html":
                self._export_summary_html_md(text_md, target_path)
            elif fmt == "md":
                with open(target_path, "w", encoding="utf-8") as f:
                    f.write(text_md)
            else:
                with open(target_path, "w", encoding="utf-8") as f:
                    f.write(markdown_to_plain(text_md))
            log.info("Summary экспортирован (%s): %s", fmt, target_path)
            QMessageBox.information(
                self, "Экспорт summary",
                f"Файл сохранён:\n{target_path}",
            )
        except Exception as exc:
            log.exception("Ошибка экспорта summary: %s", exc)
            QMessageBox.critical(
                self, "Экспорт summary",
                f"Не удалось сохранить: {exc}",
            )

    @staticmethod
    def _export_summary_html_md(text_md: str, path: str) -> None:
        from PySide6.QtGui import QTextDocument
        doc = QTextDocument()
        doc.setMarkdown(text_md)
        html_body = doc.toHtml()
        with open(path, "w", encoding="utf-8") as f:
            f.write(html_body)

    @staticmethod
    def _safe_name(name: str) -> str:
        bad = '<>:"/\\|?*\n\r\t'
        cleaned = "".join(("_" if c in bad else c)
                          for c in (name or "summary"))
        cleaned = cleaned.strip() or "summary"
        return cleaned[:60]

    # ------------------------------------------------------------------
    # Быстрое сохранение промпта в «Загрузки»
    # ------------------------------------------------------------------
    @staticmethod
    def _downloads_dir() -> str:
        try:
            import subprocess
            out = subprocess.check_output(
                ["xdg-user-dir", "DOWNLOAD"],
                text=True, timeout=2,
            ).strip()
            if out and os.path.isdir(out):
                return out
        except Exception:
            pass

        home = os.path.expanduser("~")
        for candidate in ("Загрузки", "Downloads"):
            path = os.path.join(home, candidate)
            if os.path.isdir(path):
                return path
        return home

    def _save_prompt_to_downloads(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        candidates = []
        for fname in ("deepseek_prompt.docx",
                      "deepseek_prompt.md",
                      "deepseek_prompt.txt"):
            p = os.path.join(r["dir"], fname)
            if os.path.exists(p):
                candidates.append(p)

        if not candidates:
            QMessageBox.information(
                self, "DeepSeek",
                "Для этой записи промпт не сформирован.\n\n"
                "Промпт создаётся при обработке, если в метаданных записи "
                "включён флаг «Сформировать файл промпта для DeepSeek».\n\n"
                "Если флаг не стоял — используйте «Очередь → "
                "Редактировать метаданные и перезапустить…» (Ctrl+E).",
            )
            return

        downloads = self._downloads_dir()
        if not os.path.isdir(downloads):
            QMessageBox.warning(
                self, "DeepSeek",
                f"Папка «Загрузки» не найдена:\n{downloads}",
            )
            return

        safe_name = self._safe_name(r["name"]) or "record"
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

        saved: List[str] = []
        errors: List[str] = []

        for src in candidates:
            ext = os.path.splitext(src)[1]
            base = os.path.basename(src)
            stem = os.path.splitext(base)[0]
            dst_name = f"{stem}_{safe_name}_{stamp}{ext}"
            dst = os.path.join(downloads, dst_name)

            if os.path.exists(dst):
                i = 1
                while os.path.exists(dst):
                    dst = os.path.join(
                        downloads,
                        f"{stem}_{safe_name}_{stamp}_{i}{ext}",
                    )
                    i += 1

            try:
                shutil.copy2(src, dst)
                saved.append(dst)
                log.info("Промпт сохранён в Загрузки: %s → %s", src, dst)
            except Exception as exc:
                log.exception("Ошибка сохранения промпта %s: %s", src, exc)
                errors.append(f"{base}: {exc}")

        if saved and not errors:
            lines = "\n".join(saved)
            QMessageBox.information(
                self, "DeepSeek",
                f"Файлы промпта сохранены в «Загрузки»:\n\n{lines}",
            )
        elif saved and errors:
            QMessageBox.warning(
                self, "DeepSeek",
                "Часть файлов сохранена, часть — с ошибкой:\n\n"
                + "Сохранено:\n"
                + "\n".join(saved)
                + "\n\nОшибки:\n"
                + "\n".join(errors),
            )
        else:
            QMessageBox.critical(
                self, "DeepSeek",
                "Не удалось сохранить ни одного файла:\n\n"
                + "\n".join(errors),
            )

    # ------------------------------------------------------------------
    # Утилиты: работа с видеофайлами
    # ------------------------------------------------------------------
    @staticmethod
    def _format_size(size_bytes: int) -> str:
        size = float(size_bytes)
        for unit in ("Б", "КБ", "МБ", "ГБ", "ТБ"):
            if size < 1024:
                return f"{size:.1f} {unit}"
            size /= 1024
        return f"{size:.1f} ПБ"

    def _iter_session_dirs(self) -> List[str]:
        if not os.path.isdir(self.sessions_root):
            return []
        try:
            entries = sorted(os.listdir(self.sessions_root))
        except OSError as exc:
            log.error("Не удалось прочитать %s: %s", self.sessions_root, exc)
            return []

        result: List[str] = []
        for name in entries:
            d = os.path.join(self.sessions_root, name)
            if os.path.isdir(d):
                result.append(d)
        return result

    @staticmethod
    def _find_video_file(session_dir: str) -> str:
        for ext in _VIDEO_EXTS:
            p = os.path.join(session_dir, f"video{ext}")
            if os.path.exists(p):
                return p
        return ""

    @staticmethod
    def _find_audio_file(session_dir: str) -> str:
        for ext in _AUDIO_EXTS:
            p = os.path.join(session_dir, f"video{ext}")
            if os.path.exists(p):
                return p
        return ""

    def _collect_video_stats(self) -> Dict[str, Any]:
        result: Dict[str, Any] = {
            "sessions_total": 0,
            "sessions_with_video": 0,
            "sessions_without_audio": 0,
            "deletable": [],
            "total_size": 0,
            "keep_size": 0,
        }

        for session_dir in self._iter_session_dirs():
            result["sessions_total"] += 1

            video_path = self._find_video_file(session_dir)
            if not video_path:
                continue

            ext = os.path.splitext(video_path)[1].lower()
            if ext in _AUDIO_EXTS:
                continue

            result["sessions_with_video"] += 1

            try:
                size = os.path.getsize(video_path)
            except OSError:
                size = 0

            audio_path = self._find_audio_file(session_dir)
            if not audio_path:
                result["sessions_without_audio"] += 1
                result["keep_size"] += size
                continue

            result["deletable"].append({
                "session": session_dir,
                "name": os.path.basename(session_dir),
                "video": video_path,
                "video_size": size,
                "audio": audio_path,
            })
            result["total_size"] += size

        return result

    def _show_video_stats(self) -> None:
        stats = self._collect_video_stats()

        total = stats["sessions_total"]
        with_video = stats["sessions_with_video"]
        without_audio = stats["sessions_without_audio"]
        deletable = len(stats["deletable"])
        total_size = self._format_size(stats["total_size"])

        msg = (
            f"<b>Всего сессий:</b> {total}<br>"
            f"<b>Сессий с video.mp4:</b> {with_video}<br>"
            f"<b>Сессий без аудио (удалять нельзя):</b> {without_audio}<br>"
            f"<br>"
            f"<b>Можно удалить video.mp4:</b> {deletable} шт.<br>"
            f"<b>Освободится:</b> {total_size}"
        )

        QMessageBox.information(self, "Размер видеофайлов", msg)
        log.info(
            "Статистика видеофайлов: всего сессий=%d, с видео=%d, "
            "без аудио=%d, можно удалить=%d (%s)",
            total, with_video, without_audio, deletable, total_size,
        )

    def _delete_video_files(self) -> None:
        stats = self._collect_video_stats()
        deletable: List[Dict[str, Any]] = stats["deletable"]

        if not deletable:
            QMessageBox.information(
                self, "Удаление видеофайлов",
                "Нет сессий, где можно удалить video.mp4.\n\n"
                "Удаляются только сессии, в которых уже есть аудиофайл "
                "(video.mp3 / .aac / .wav / .opus / .ogg / .m4a).\n"
                "Сессии без аудио не трогаются, чтобы не потерять запись.",
            )
            return

        total_size_str = self._format_size(stats["total_size"])

        preview_lines: List[str] = []
        preview_limit = 15
        for i, item in enumerate(deletable):
            if i >= preview_limit:
                preview_lines.append(
                    f"… и ещё {len(deletable) - preview_limit} сессий"
                )
                break
            size_str = self._format_size(item["video_size"])
            preview_lines.append(f"• {item['name']} — {size_str}")

        preview = "\n".join(preview_lines)

        msg = (
            f"<b>Будет удалён video.mp4 из сессий:</b> {len(deletable)} шт.<br>"
            f"<b>Освободится места:</b> {total_size_str}<br>"
            f"<br>"
            f"Аудиофайлы (video.mp3 и др.) <b>останутся</b>.<br>"
            f"Сессии без аудио не трогаются.<br>"
            f"<br>"
            f"<b>Сессии к удалению:</b><br>"
            f"<pre style='font-family:monospace'>{preview}</pre>"
            f"<br>"
            f"<b style='color:#c62828'>Действие необратимо.</b><br>"
            f"Продолжить?"
        )

        reply = QMessageBox.question(
            self,
            "Подтверждение удаления",
            msg,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            log.info("Удаление видеофайлов отменено пользователем")
            return

        progress = QProgressDialog(
            "Удаление видеофайлов…",
            "Отмена",
            0, len(deletable),
            self,
        )
        progress.setWindowTitle("Удаление видеофайлов")
        progress.setWindowModality(Qt.WindowModality.WindowModal)
        progress.setMinimumDuration(0)
        progress.setValue(0)

        removed = 0
        freed_bytes = 0
        errors: List[str] = []

        for i, item in enumerate(deletable):
            if progress.wasCanceled():
                log.warning("Удаление видеофайлов отменено пользователем "
                            "на шаге %d/%d", i, len(deletable))
                break

            video_path = item["video"]
            size = item["video_size"]

            audio_path = self._find_audio_file(item["session"])
            if not audio_path:
                errors.append(
                    f"{item['name']}: аудио исчезло, пропуск"
                )
                log.warning("Пропуск %s: аудио исчезло перед удалением",
                            video_path)
                progress.setValue(i + 1)
                continue

            try:
                os.remove(video_path)
                removed += 1
                freed_bytes += size
                log.info("Удалён видеофайл: %s (%s)",
                         video_path, self._format_size(size))
            except Exception as exc:
                errors.append(f"{item['name']}: {exc}")
                log.exception("Не удалось удалить %s: %s", video_path, exc)

            progress.setValue(i + 1)

        progress.setValue(len(deletable))

        freed_str = self._format_size(freed_bytes)

        if errors:
            QMessageBox.warning(
                self, "Удаление видеофайлов",
                f"Удалено файлов: {removed} из {len(deletable)}\n"
                f"Освобождено: {freed_str}\n\n"
                f"Ошибки при удалении:\n"
                + "\n".join(errors[:20])
                + ("\n…" if len(errors) > 20 else ""),
            )
        else:
            QMessageBox.information(
                self, "Удаление видеофайлов",
                f"Удалено файлов: {removed}\n"
                f"Освобождено: {freed_str}",
            )

        log.info(
            "Удаление видеофайлов завершено: удалено=%d, освобождено=%s, "
            "ошибок=%d",
            removed, freed_str, len(errors),
        )

        self.refresh()

    # ------------------------------------------------------------------
    # Внутренние операции
    # ------------------------------------------------------------------
    def _remove_from_queue(self, task_id: str) -> bool:
        if not task_id:
            return False
        try:
            self.task_queue.remove_task(task_id)
            log.info("Задача %s удалена из очереди", task_id)
            return True
        except Exception as exc:
            log.warning("Не удалось удалить задачу %s: %s", task_id, exc)
            return False

    def _delete_processed_files(self, session_dir: str) -> None:
        for fname in ("video.mp3", "video.txt", "video.aac",
                      "video.wav", "video.opus"):
            p = os.path.join(session_dir, fname)
            if os.path.exists(p):
                try:
                    os.remove(p)
                    log.info("Удалён файл: %s", p)
                except Exception as exc:
                    log.warning("Не удалось удалить %s: %s", p, exc)

    def _remove_processed_artifacts(self, session_dir: str) -> None:
        patterns = [
            "video.mp3", "video.aac", "video.wav", "video.opus",
            "video.ogg", "video.m4a",
            "video.txt",
            "video_summary.md",
            "deepseek_prompt.txt", "deepseek_prompt.md",
            "deepseek_prompt.docx",
            "protocol.docx", "protocol.md", "protocol.txt",
        ]
        for fname in patterns:
            p = os.path.join(session_dir, fname)
            if os.path.exists(p):
                try:
                    os.remove(p)
                    log.info("Удалён артефакт обработки: %s", p)
                except Exception as exc:
                    log.warning("Не удалось удалить %s: %s", p, exc)

    def _enqueue_session(self, r: Dict[str, Any]) -> Optional[str]:
        if not os.path.exists(r["video_path"]):
            log.error("Видео не найдено: %s", r["video_path"])
            return None
        meta = read_json_file(os.path.join(r["dir"], "session.json")) or {}
        meta["video_path"] = r["video_path"]
        meta["session_dir"] = r["dir"]
        meta.pop("task_id", None)
        try:
            task_id = self.task_queue.add_task(
                {"video_path": r["video_path"], "metadata": meta}
            )
            log.info("Задача %s добавлена в очередь (session=%s)",
                     task_id, r["dir"])
            return task_id
        except Exception as exc:
            log.exception("Не удалось добавить задачу: %s", exc)
            return None

    # ------------------------------------------------------------------
    # Перезапуск / редактирование метаданных
    # ------------------------------------------------------------------
    def _restart_processing(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        if not r["video_path"] or not os.path.exists(r["video_path"]):
            QMessageBox.warning(self, "Записи",
                                f"Видео не найдено:\n{r['video_path']}")
            return
        if QMessageBox.question(
            self, "Перезапустить обработку",
            f"Запись «{r['name']}» будет обработана заново с текущими "
            "параметрами:\n"
            "• старая задача удалится из очереди,\n"
            "• файлы video.mp3 / video.txt / deepseek_prompt.* удалятся,\n"
            "• запись добавится в очередь заново.\n\nПродолжить?",
        ) != QMessageBox.StandardButton.Yes:
            return
        if r["task_id"]:
            self._remove_from_queue(r["task_id"])
        self._remove_processed_artifacts(r["dir"])
        task_id = self._enqueue_session(r)
        if task_id:
            QMessageBox.information(self, "Записи",
                                    f"Запись добавлена в очередь: {task_id}")
            self.refresh()
        else:
            QMessageBox.critical(self, "Ошибка",
                                 "Не удалось добавить запись в очередь")

    def _edit_metadata_and_restart(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        if not r["video_path"] or not os.path.exists(r["video_path"]):
            QMessageBox.warning(
                self, "Записи",
                f"Видео не найдено:\n"
                f"{r['video_path'] or '(путь не задан)'}",
            )
            return

        if self.config_manager is None:
            QMessageBox.warning(
                self, "Записи",
                "Нет доступа к настройкам — ConfigManager не передан.",
            )
            return

        session_json = os.path.join(r["dir"], "session.json")
        meta = read_json_file(session_json) or {}

        projects = self.config_manager.get_project_names()
        prompts = self.config_manager.get_prompts()
        default_prompt = self.config_manager.get_default_prompt()
        name_templates = self.config_manager.get_name_templates()

        project = meta.get("project")
        if project and project not in projects:
            projects = [project] + projects

        dlg = MetadataDialog(
            projects=projects,
            prompts=prompts,
            title=f"Метаданные записи — {r['name']}",
            initial=meta,
            default_prompt=default_prompt,
            sessions_root=self.sessions_root,
            on_save_prompt=self._on_save_prompt_to_config,
            get_prompts=self.config_manager.get_prompts,
            name_templates=name_templates,
            on_save_name_template=self._on_save_name_template,
            get_name_templates=self.config_manager.get_name_templates,
            parent=self,
        )
        if dlg.exec() != QDialog.DialogCode.Accepted:
            log.info("Редактирование метаданных отменено пользователем")
            return

        new_meta = dlg.result_data

        for keep_key in (
            "date", "time", "monitor",
            "summary_bb",
            "manual_protocol_path",
            "source", "source_files",
            "video_path", "session_dir",
        ):
            if keep_key not in new_meta and keep_key in meta:
                new_meta[keep_key] = meta[keep_key]

        new_meta.setdefault("video_path", r["video_path"])
        new_meta.setdefault("session_dir", r["dir"])
        new_meta.pop("task_id", None)

        reply = QMessageBox.question(
            self,
            "Перезапустить обработку",
            f"Запись «{r['name']}» будет обработана заново с новыми "
            f"параметрами:<br><br>"
            f"• скрам-митинг: "
            f"<b>{'да' if new_meta.get('is_scrum') else 'нет'}</b><br>"
            f"• формировать summary: "
            f"<b>{'да' if new_meta.get('generate_summary') else 'нет'}</b>"
            f"<br>"
            f"• сформировать промпт DeepSeek: "
            f"<b>{'да' if new_meta.get('generate_deepseek_prompt') else 'нет'}</b>"
            f"<br><br>"
            f"Старые файлы обработки (video.txt, video.mp3, "
            f"deepseek_prompt.*) будут удалены.<br><br>"
            f"Продолжить?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.Yes,
        )
        if reply != QMessageBox.StandardButton.Yes:
            log.info("Перезапуск с новыми метаданными отменён")
            return

        if not self._write_json(session_json, new_meta):
            QMessageBox.critical(
                self, "Записи",
                "Не удалось сохранить session.json",
            )
            return
        log.info(
            "Метаданные записи обновлены: %s (is_scrum=%s, "
            "generate_summary=%s, generate_deepseek=%s)",
            r["dir"],
            new_meta.get("is_scrum"),
            new_meta.get("generate_summary"),
            new_meta.get("generate_deepseek_prompt"),
        )

        self._remove_processed_artifacts(r["dir"])

        if r["task_id"]:
            self._remove_from_queue(r["task_id"])

        task_id = self._enqueue_session({
            "video_path": r["video_path"],
            "dir": r["dir"],
        })
        if not task_id:
            QMessageBox.critical(
                self, "Записи",
                "Метаданные сохранены, но не удалось поставить запись "
                "в очередь. Проверьте лог.",
            )
            return

        QMessageBox.information(
            self, "Записи",
            f"Запись поставлена в очередь на обработку: {task_id}",
        )
        self.refresh()

    # ------------------------------------------------------------------
    # Помощники для MetadataDialog
    # ------------------------------------------------------------------
    def _on_save_prompt_to_config(self, name: str, text: str) -> None:
        try:
            name = (name or "").strip()
            text = (text or "").strip()
            if not name or not text:
                return
            cfg = self.config_manager.config
            meta = cfg.setdefault("metadata", {})
            prompts = meta.setdefault("prompts", [])
            replaced = False
            for i, p in enumerate(prompts):
                if isinstance(p, dict) and p.get("name") == name:
                    prompts[i] = {"name": name, "text": text}
                    replaced = True
                    break
            if not replaced:
                prompts.append({"name": name, "text": text})
            self.config_manager.save()
            log.info(
                "Промпт «%s» %s в библиотеку (%d символов)",
                name, "обновлён" if replaced else "добавлен", len(text),
            )
        except Exception as exc:
            log.exception("Ошибка сохранения промпта: %s", exc)
            raise

    def _on_save_name_template(self, label: str, template: str) -> None:
        try:
            label = (label or "").strip()
            template = (template or "").strip()
            if not template:
                return
            self.config_manager.add_name_template(label, template)
            log.info("Шаблон имени сохранён: label=%r, template=%r",
                     label, template)
        except Exception as exc:
            log.exception("Ошибка сохранения шаблона имени: %s", exc)
            raise

    # ------------------------------------------------------------------
    # Изменение статуса
    # ------------------------------------------------------------------
    def _change_status(self, new_status: str) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        confirm_msg = {
            STATUS_UPLOADED:
                f"Сбросить статус «{r['name']}» в «Сохранено»?\n\n"
                "• задача удалится из очереди,\n"
                "• файлы video.mp3 / video.txt удалятся.",
            STATUS_PROCESSED:
                f"Пометить «{r['name']}» как «Обработан»?\n\n"
                "• задача удалится из очереди,\n"
                "• файлы video.mp3 / video.txt останутся.",
            STATUS_ERROR:
                f"Пометить «{r['name']}» как «Ошибка»?\n\n"
                "• задача удалится из очереди,\n"
                "• файлы видео и транскрипта останутся.",
        }.get(new_status, f"Изменить статус на «{new_status}»?")
        if QMessageBox.question(self, "Изменить статус", confirm_msg) != \
                QMessageBox.StandardButton.Yes:
            return
        if r["task_id"]:
            self._remove_from_queue(r["task_id"])
        if new_status == STATUS_UPLOADED:
            self._delete_processed_files(r["dir"])
        if new_status == STATUS_PROCESSED:
            has_audio = os.path.exists(
                os.path.join(r["dir"], "video.mp3")
            )
            has_txt = os.path.exists(
                os.path.join(r["dir"], "video.txt")
            )
            if not (has_audio or has_txt):
                QMessageBox.warning(
                    self, "Записи",
                    "Не найдены video.mp3 или video.txt — файлы "
                    "обработки отсутствуют.\n\nСтатус изменён не будет. "
                    "Используйте «Перезапустить», если нужно обработать "
                    "запись.",
                )
                self.refresh()
                return
        log.info("Ручное изменение статуса %s: %s → %s",
                 r["dir"], r["status"], new_status)
        self.refresh()

    def _enqueue_current(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        if r["task_id"]:
            QMessageBox.information(self, "Записи",
                                    f"Запись уже в очереди: {r['task_id']}")
            return
        if not r["video_path"] or not os.path.exists(r["video_path"]):
            QMessageBox.warning(self, "Записи",
                                f"Видео не найдено:\n{r['video_path']}")
            return
        task_id = self._enqueue_session(r)
        if task_id:
            QMessageBox.information(self, "Записи",
                                    f"Запись добавлена в очередь: {task_id}")
            self.refresh()
        else:
            QMessageBox.critical(self, "Ошибка",
                                 "Не удалось добавить запись в очередь")

    # ------------------------------------------------------------------
    # Контекстное меню по правому клику
    # ------------------------------------------------------------------
    def _show_context_menu(self, pos) -> None:
        r = self._selected_row()
        if not r:
            return
        menu = QMenu(self)
        menu.addAction(
            "Редактировать метаданные и перезапустить…",
            self._edit_metadata_and_restart,
        )
        menu.addSeparator()
        menu.addAction("Перезапустить обработку", self._restart_processing)
        menu.addAction("Поставить в очередь", self._enqueue_current)
        menu.addSeparator()
        menu.addAction("Открыть папку записи", self._open_folder)
        menu.addAction("Открыть видео", self._open_video)
        menu.addSeparator()
        menu.addAction(
            "Изменить summary (Markdown)…", self._edit_summary_bb
        )
        menu.addAction(
            "Создать/редактировать протокол (Markdown)…",
            self._edit_manual_protocol_md,
        )
        menu.exec(self.table.viewport().mapToGlobal(pos))

    # ------------------------------------------------------------------
    # DeepSeek-промпт
    # ------------------------------------------------------------------
    def _open_deepseek_prompt(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        path = r.get("prompt_path", "")
        if not path or not os.path.exists(path):
            QMessageBox.information(
                self, "DeepSeek",
                "Для этой записи промпт не сформирован.\n\n"
                "Промпт создаётся при обработке, если в метаданных записи "
                "включён флаг «Сформировать файл промпта для DeepSeek».\n\n"
                "Если флаг не стоял — используйте «Очередь → "
                "Редактировать метаданные и перезапустить…» (Ctrl+E).",
            )
            return
        log.info("Открытие DeepSeek-промпта: %s", path)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def _export_prompt(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        path = r.get("prompt_path", "")
        if not path or not os.path.exists(path):
            QMessageBox.information(
                self, "DeepSeek",
                "Для этой записи промпт не сформирован.\n\n"
                "Если флаг не стоял — используйте «Очередь → "
                "Редактировать метаданные и перезапустить…» (Ctrl+E).",
            )
            return

        preferred = "docx"
        if self.config_manager is not None:
            try:
                preferred = self.config_manager.get_scrum_settings().get(
                    "export_format", "docx"
                )
            except Exception:
                pass

        formats = ["docx", "md", "txt"]
        try:
            pref_idx = formats.index(preferred)
        except ValueError:
            pref_idx = 0

        fmt, ok = QInputDialog.getItem(
            self, "Экспорт промпта", "Формат файла:",
            formats, pref_idx, False,
        )
        if not ok or not fmt:
            return

        text = read_any_text(path)
        if not text.strip():
            QMessageBox.warning(
                self, "Экспорт",
                "Не удалось прочитать промпт.",
            )
            return

        base_name = os.path.splitext(os.path.basename(path))[0]
        default_path = os.path.join(r["dir"], f"{base_name}.{fmt}")
        target_path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить промпт как",
            default_path,
            "Word (*.docx);;Markdown (*.md);;Text (*.txt);;All files (*)",
        )
        if not target_path:
            return

        try:
            if target_path.endswith(".docx"):
                from docx import Document
                doc = Document()
                for line in text.splitlines():
                    doc.add_paragraph(line)
                doc.save(target_path)
            else:
                with open(target_path, "w", encoding="utf-8") as f:
                    f.write(text)
            log.info("Промпт экспортирован: %s", target_path)
            QMessageBox.information(self, "Экспорт",
                                    f"Файл сохранён:\n{target_path}")
        except Exception as exc:
            log.exception("Ошибка сохранения промпта: %s", exc)
            QMessageBox.critical(self, "Экспорт",
                                 f"Не удалось сохранить: {exc}")

    # ------------------------------------------------------------------
    # Вложения
    # ------------------------------------------------------------------
    def _open_attachments(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        att_dir = os.path.join(r["dir"], "attachments")
        if not os.path.isdir(att_dir):
            QMessageBox.information(
                self, "Вложения",
                "У этой записи нет вложений.",
            )
            return
        log.info("Открытие папки вложений: %s", att_dir)
        QDesktopServices.openUrl(QUrl.fromLocalFile(att_dir))

    def _add_attachment_to_session(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return

        files, _ = QFileDialog.getOpenFileNames(
            self,
            "Выберите файлы-вложения",
            os.path.expanduser("~"),
            "Документы (*.txt *.md *.docx *.pdf *.csv *.json);;"
            "Все файлы (*)",
        )
        if not files:
            return

        att_dir = os.path.join(r["dir"], "attachments")
        os.makedirs(att_dir, exist_ok=True)

        session_json = os.path.join(r["dir"], "session.json")
        meta = read_json_file(session_json) or {}
        current = list(meta.get("attachments", []) or [])

        added = 0
        for src in files:
            if not os.path.exists(src):
                continue
            base = os.path.basename(src)
            dst = os.path.join(att_dir, base)
            if os.path.exists(dst):
                stem, ext = os.path.splitext(base)
                i = 1
                while os.path.exists(dst):
                    dst = os.path.join(att_dir, f"{stem}_{i}{ext}")
                    i += 1
            try:
                shutil.copy2(src, dst)
                current.append(dst)
                added += 1
                log.info("Добавлено вложение: %s → %s", src, dst)
            except Exception as exc:
                log.exception("Ошибка копирования вложения %s: %s",
                              src, exc)

        if added:
            meta["attachments"] = current
            if not self._write_json(session_json, meta):
                QMessageBox.critical(self, "Ошибка",
                                     "Не удалось обновить session.json")
                return
            QMessageBox.information(
                self, "Вложения",
                f"Добавлено файлов: {added}\n\n"
                "Чтобы вложения попали в транскрибацию или промпт "
                "DeepSeek, установите соответствующие флаги в "
                "метаданных через «Очередь → Редактировать метаданные "
                "и перезапустить…».",
            )
            self.refresh()

    # ------------------------------------------------------------------
    # Обычные действия
    # ------------------------------------------------------------------
    def _selected_row(self) -> Optional[Dict[str, Any]]:
        i = self.table.currentRow()
        if i < 0 or i >= len(self._rows):
            return None
        return self._rows[i]

    def _open_folder(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        folder = r["dir"]
        if not os.path.isdir(folder):
            QMessageBox.warning(self, "Записи",
                                f"Папка не найдена:\n{folder}")
            return
        log.info("Открытие папки сессии: %s", folder)
        QDesktopServices.openUrl(QUrl.fromLocalFile(folder))

    def _open_video(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        video = r["video_path"]
        if not video or not os.path.exists(video):
            QMessageBox.warning(self, "Записи",
                                f"Видео не найдено:\n{video}")
            return
        log.info("Открытие видео: %s", video)
        QDesktopServices.openUrl(QUrl.fromLocalFile(video))

    def _delete_session(self) -> None:
        r = self._selected_row()
        if not r:
            QMessageBox.warning(self, "Записи", "Выберите запись")
            return
        if QMessageBox.question(
            self, "Удалить запись",
            f"Удалить папку записи «{r['name']}» со всем содержимым?\n\n"
            f"{r['dir']}\n\nДействие необратимо.",
        ) != QMessageBox.StandardButton.Yes:
            return
        if r["task_id"]:
            self._remove_from_queue(r["task_id"])
        try:
            shutil.rmtree(r["dir"])
            log.warning("Папка сессии удалена: %s", r["dir"])
            self.refresh()
        except Exception as exc:
            log.exception("Не удалось удалить %s: %s", r["dir"], exc)
            QMessageBox.critical(self, "Ошибка",
                                 f"Не удалось удалить: {exc}")

    # ------------------------------------------------------------------
    # Справка
    # ------------------------------------------------------------------
    def _show_shortcuts(self) -> None:
        QMessageBox.information(
            self,
            "Горячие клавиши",
            "Ctrl+I        — импорт материалов\n"
            "Ctrl+M        — создать/редактировать протокол (Markdown)\n"
            "Ctrl+Shift+M  — экспорт протокола в DOCX\n"
            "Ctrl+B        — отправить протокол/summary в Bitrix24\n"
            "Ctrl+E        — редактировать метаданные и перезапустить\n"
            "Ctrl+Shift+E  — открыть папку записи\n"
            "Ctrl+Shift+V  — открыть видео\n"
            "Ctrl+P        — изменить summary (Markdown)\n"
            "Ctrl+Shift+P  — просмотр summary\n"
            "Ctrl+D        — открыть промпт DeepSeek\n"
            "Ctrl+Shift+D  — экспорт промпта\n"
            "Ctrl+Alt+D    — сохранить промпт в «Загрузки»\n"
            "Ctrl+R        — перезапустить обработку\n"
            "F5            — обновить список\n"
            "Ctrl+Delete   — удалить запись\n"
            "Ctrl+W        — закрыть окно\n"
            "\n"
            "Утилиты → «Удалить видеофайлы» — освободить место, "
            "оставив только аудио (с подтверждением).",
        )