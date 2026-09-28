"""Конвертер Markdown → DOCX для протоколов.

Использует только python-docx (уже есть в зависимостях проекта).
Поддерживает подмножество Markdown, достаточное для протоколов:

    # H1, ## H2, ### H3 (и до ######)
    **жирный**, *курсив*, `инлайн-код`
    - маркированный список
    1. нумерованный список
    > цитата
    --- горизонтальная линия
    ```блок кода```
    [текст](url) — url добавляется в скобках после текста

Никакой HTML-санитайзации не делает — текст считается доверенным.
"""
from __future__ import annotations

import re
from typing import List, Optional

from .logger import get_logger

log = get_logger(__name__)


# ---------------------------------------------------------------------------
# Разбор инлайн-разметки
# ---------------------------------------------------------------------------
_INLINE_TOKEN_RE = re.compile(
    r"(\*\*.+?\*\*"          # **bold**
    r"|\*[^*\n]+?\*"         # *italic*
    r"|`[^`\n]+?`"           # `code`
    r"|\[[^\]]+?\]\([^)]+?\)"  # [text](url)
    r")"
)


def _add_inline(paragraph, text: str) -> None:
    """Разбирает инлайн-разметку строки и добавляет runs в paragraph."""
    # Сначала упрощаем: экранируем ничего не надо, работаем как есть.
    parts = _INLINE_TOKEN_RE.split(text)
    for part in parts:
        if not part:
            continue

        # **bold**
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            continue

        # *italic* (не путать с **bold** — порядок проверок важен)
        if (
            part.startswith("*")
            and part.endswith("*")
            and len(part) > 2
            and not part.startswith("**")
        ):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            continue

        # `code`
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            continue

        # [text](url)
        m = re.match(r"\[([^\]]+?)\]\(([^)]+?)\)$", part)
        if m:
            run = paragraph.add_run(m.group(1))
            run.italic = False
            # URL добавляем в скобках после текста — как в обычном
            # текстовом представлении, чтобы при копировании из .docx
            # он не терялся.
            tail = paragraph.add_run(f" ({m.group(2)})")
            tail.font.color.rgb = None  # оставляем как есть
            continue

        # Обычный текст
        paragraph.add_run(part)


# ---------------------------------------------------------------------------
# Разбор блочной структуры
# ---------------------------------------------------------------------------
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_ULIST_RE = re.compile(r"^[-*+]\s+(.*)$")
_OLIST_RE = re.compile(r"^(\d+)\.\s+(.*)$")
_HR_RE = re.compile(r"^-{3,}$|^\*{3,}$|^_{3,}$")
_CODE_FENCE_RE = re.compile(r"^```")


def markdown_to_docx(md_text: str, output_path: str, title: str = "") -> str:
    """
    Конвертирует Markdown-текст в .docx и сохраняет по пути output_path.

    Args:
        md_text:     исходный Markdown.
        output_path: куда сохранить .docx.
        title:       если задан — добавляется как заголовок H1 в начало
                     (если в md_text нет своего H1).

    Returns:
        Путь к сохранённому файлу.

    Raises:
        RuntimeError: если python-docx не установлен или сохранение
                      не удалось.
    """
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "python-docx не установлен. Установите: pip install python-docx"
        ) from exc

    doc = Document()

    # Базовые стили — сделаем шрифт чуть приятнее для протоколов.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = (md_text or "").splitlines()
    i = 0
    n = len(lines)

    # Заголовок по умолчанию — если в тексте нет своего H1.
    has_h1 = any(_HEADING_RE.match(l) and len(_HEADING_RE.match(l).group(1)) == 1
                 for l in lines if l.strip())
    if title and not has_h1:
        doc.add_heading(title, level=1)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # --- Пустая строка ---
        if not stripped:
            i += 1
            continue

        # --- Блок кода ---
        if _CODE_FENCE_RE.match(stripped):
            i += 1
            code_lines: List[str] = []
            while i < n and not _CODE_FENCE_RE.match(lines[i].strip()):
                code_lines.append(lines[i])
                i += 1
            # пропускаем закрывающие ```
            if i < n:
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            p.paragraph_format.left_indent = Pt(12)
            continue

        # --- Горизонтальная линия ---
        if _HR_RE.match(stripped):
            p = doc.add_paragraph()
            p.add_run("─" * 40).font.color.rgb = None
            i += 1
            continue

        # --- Заголовки ---
        m = _HEADING_RE.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)  # python-docx: Heading 1..9
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        # --- Цитата ---
        if stripped.startswith(">"):
            quote_lines: List[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            text = " ".join(quote_lines).strip()
            p = doc.add_paragraph()
            _add_inline(p, text)
            p.paragraph_format.left_indent = Pt(24)
            for run in p.runs:
                run.italic = True
            continue

        # --- Маркированный список ---
        m = _ULIST_RE.match(stripped)
        if m:
            while i < n:
                m2 = _ULIST_RE.match(lines[i].strip())
                if not m2:
                    break
                p = doc.add_paragraph(style="List Bullet")
                _add_inline(p, m2.group(1))
                i += 1
            continue

        # --- Нумерованный список ---
        m = _OLIST_RE.match(stripped)
        if m:
            while i < n:
                m2 = _OLIST_RE.match(lines[i].strip())
                if not m2:
                    break
                p = doc.add_paragraph(style="List Number")
                _add_inline(p, m2.group(2))
                i += 1
            continue

        # --- Обычный абзац ---
        p = doc.add_paragraph()
        _add_inline(p, stripped)
        i += 1

    try:
        doc.save(output_path)
    except Exception as exc:
        raise RuntimeError(f"Не удалось сохранить DOCX: {exc}") from exc

    log.info("Markdown → DOCX сохранён: %s", output_path)
    return output_path