"""Формирование файла-промпта для извлечения поручений.

Модуль не зависит от Qt и от VideoProcessor — это чистая функция,
которую можно вызывать:
  • из processor.py во время обработки записи;
  • из UI (кнопка «Обновить промпт поручений») после ручной
    правки протокола;
  • из CLI/скриптов.

Что делает:
  • находит протокол в папке сессии (manual_protocol.* / protocol.*);
  • читает его через file_readers.read_any_text (умеет .docx/.pdf/.md/.txt);
  • собирает документ: инструкция + метаданные + протокол;
  • сохраняет в action_items_prompt.<ext> (docx/md/txt).
"""
from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from .file_readers import read_any_text
from .logger import get_logger
from .config_manager import DEFAULT_ACTION_ITEMS_PROMPT

log = get_logger(__name__)


# Порядок приоритетов при поиске протокола в папке сессии.
# Сначала ручной протокол, затем обычный.
PROTOCOL_CANDIDATES: List[str] = [
    "manual_protocol.docx",
    "manual_protocol.md",
    "manual_protocol.txt",
    "manual_protocol.pdf",
    "protocol.docx",
    "protocol.md",
    "protocol.txt",
    "protocol.pdf",
]


# Предел чтения протокола — чтобы случайно не затянуть
# гигантский PDF в память.
DEFAULT_MAX_CHARS = 5_000_000


# ---------------------------------------------------------------------------
# Публичный API
# ---------------------------------------------------------------------------
def find_protocol_in_session(session_dir: str) -> str:
    """
    Возвращает путь к найденному протоколу в папке сессии.

    Порядок приоритетов — PROTOCOL_CANDIDATES.
    Возвращает "" если ничего не найдено.
    """
    if not session_dir or not os.path.isdir(session_dir):
        return ""
    for fname in PROTOCOL_CANDIDATES:
        p = os.path.join(session_dir, fname)
        if os.path.isfile(p):
            return p
    return ""


def read_protocol_text(
    session_dir: str,
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> tuple[str, str]:
    """
    Ищет протокол в папке сессии и читает его как plain text.

    Returns:
        (текст, путь_к_протоколу)
        Если протокола нет — ("", "").
    """
    path = find_protocol_in_session(session_dir)
    if not path:
        return "", ""

    try:
        text = read_any_text(path, max_chars)
    except Exception as exc:
        log.warning(
            "Не удалось прочитать протокол %s: %s", path, exc
        )
        return "", path

    if not text or not text.strip():
        log.warning(
            "Протокол %s прочитан, но пустой", path
        )
        return "", path

    return text, path


def build_prompt_blocks(
    *,
    protocol_text: str,
    session_name: str = "",
    session_date: str = "",
    template: str = "",
) -> List[Dict[str, Any]]:
    """
    Собирает список блоков документа (без записи в файл).

    Используется процессором и UI.
    """
    if not template:
        template = DEFAULT_ACTION_ITEMS_PROMPT

    blocks: List[Dict[str, Any]] = []

    # 1. Инструкция для модели.
    blocks.append({
        "kind": "text",
        "title": "ИНСТРУКЦИЯ ДЛЯ МОДЕЛИ",
        "text": template.strip(),
    })

    # 2. Метаданные сессии.
    meta_lines: List[str] = []
    if session_name:
        meta_lines.append(f"Название записи: {session_name}")
    if session_date:
        meta_lines.append(f"Дата совещания: {session_date}")
    if meta_lines:
        blocks.append({
            "kind": "text",
            "title": "ДАННЫЕ СОВЕЩАНИЯ",
            "text": "\n".join(meta_lines),
        })

    # 3. Протокол.
    if protocol_text and protocol_text.strip():
        blocks.append({
            "kind": "text",
            "title": "ПРОТОКОЛ",
            "text": protocol_text.strip(),
        })
    else:
        blocks.append({
            "kind": "text",
            "title": "ПРОТОКОЛ",
            "text": (
                "(Протокол не приложен к записи. "
                "Добавьте manual_protocol.docx "
                "или protocol.docx.)"
            ),
        })

    return blocks


def write_prompt_file(
    *,
    session_dir: str,
    blocks: List[Dict[str, Any]],
    fmt: str = "docx",
) -> str:
    """
    Сохраняет блоки в файл action_items_prompt.<ext>.

    Args:
        session_dir: папка сессии.
        blocks:      список блоков (из build_prompt_blocks).
        fmt:         "docx" | "md" | "txt".

    Returns:
        Путь к сохранённому файлу или "" при ошибке.
    """
    if not session_dir or not os.path.isdir(session_dir):
        log.warning(
            "write_prompt_file: папка не найдена: %s", session_dir
        )
        return ""

    fmt = (fmt or "docx").lower()
    if fmt not in ("docx", "md", "txt"):
        fmt = "docx"

    base_name = "action_items_prompt"

    # ---------------- TXT ----------------
    if fmt == "txt":
        path = os.path.join(session_dir, f"{base_name}.txt")
        parts: List[str] = []
        for idx, b in enumerate(blocks):
            title = (b.get("title") or "").strip()
            if idx > 0:
                parts.append("")
                parts.append("=" * 60)
                parts.append("")
            if title:
                parts.append(title)
                parts.append("=" * 60)
                parts.append("")
            parts.append(b.get("text") or "")
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write("\n".join(parts))
            log.info("Промпт поручений сохранён (txt): %s", path)
            return path
        except Exception as exc:
            log.exception(
                "Не удалось сохранить txt %s: %s", path, exc
            )
            return ""

    # ---------------- MD ----------------
    if fmt == "md":
        path = os.path.join(session_dir, f"{base_name}.md")
        md_parts: List[str] = []
        for idx, b in enumerate(blocks):
            title = (b.get("title") or "").strip()
            if idx > 0:
                md_parts.append("")
                md_parts.append("---")
                md_parts.append("")
            if title:
                md_parts.append(f"## {title}")
                md_parts.append("")
            md_parts.append(b.get("text") or "")
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write("\n".join(md_parts))
            log.info("Промпт поручений сохранён (md): %s", path)
            return path
        except Exception as exc:
            log.exception(
                "Не удалось сохранить md %s: %s", path, exc
            )
            return ""

    # ---------------- DOCX ----------------
    path = os.path.join(session_dir, f"{base_name}.docx")
    try:
        from docx import Document  # type: ignore
    except ImportError:
        log.warning(
            "python-docx не установлен — сохраняю .txt вместо .docx"
        )
        return write_prompt_file(
            session_dir=session_dir, blocks=blocks, fmt="txt"
        )

    try:
        doc = Document()

        # Базовый шрифт.
        try:
            style = doc.styles["Normal"]
            style.font.name = "Calibri"
        except Exception:
            pass

        for idx, b in enumerate(blocks):
            title = (b.get("title") or "").strip()

            if idx > 0:
                sep = doc.add_paragraph()
                run = sep.add_run("─" * 60)
                try:
                    from docx.shared import Pt, RGBColor
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                except Exception:
                    pass

            if title:
                doc.add_heading(title, level=2)

            text = b.get("text") or ""
            if text.strip():
                _append_text_to_docx(doc, text)

        doc.save(path)
        log.info(
            "Промпт поручений сохранён (docx): %s (блоков: %d)",
            path, len(blocks),
        )
        return path
    except Exception as exc:
        log.exception(
            "Не удалось сохранить docx %s: %s", path, exc
        )
        return ""


def _append_text_to_docx(doc, text: str) -> None:
    """
    Добавляет plain-текст в docx, разбивая его на абзацы.

    Никакой Markdown-разметки — только абзацы. Для служебного
    файла промпта этого достаточно.
    """
    for line in text.splitlines():
        doc.add_paragraph(line)


# ---------------------------------------------------------------------------
# Удобная «одна кнопка»
# ---------------------------------------------------------------------------
def regenerate_action_items_prompt(
    *,
    session_dir: str,
    fmt: str = "docx",
    template: str = "",
    session_name: str = "",
    session_date: str = "",
    max_chars: int = DEFAULT_MAX_CHARS,
) -> Dict[str, Any]:
    """
    Полный цикл: найти протокол → прочитать → записать prompt-файл.

    Используется UI и processor-ом как «одна функция».

    Returns:
        {
          "ok": bool,
          "path": str,             # путь к action_items_prompt.<ext>
          "protocol_path": str,    # путь к использованному протоколу
          "protocol_chars": int,
          "warning": str,
        }
    """
    result: Dict[str, Any] = {
        "ok": False,
        "path": "",
        "protocol_path": "",
        "protocol_chars": 0,
        "warning": "",
    }

    if not session_dir or not os.path.isdir(session_dir):
        result["warning"] = f"Папка сессии не найдена: {session_dir}"
        return result

    # --- Читаем протокол ---
    protocol_text, protocol_path = read_protocol_text(
        session_dir, max_chars=max_chars
    )
    result["protocol_path"] = protocol_path
    result["protocol_chars"] = len(protocol_text or "")

    if not protocol_path:
        result["warning"] = (
            "Протокол не найден в папке записи. "
            "Ожидается manual_protocol.docx / .md / .txt / .pdf "
            "или protocol.docx / .md / .txt / .pdf."
        )
    elif not protocol_text:
        result["warning"] = (
            f"Протокол найден ({os.path.basename(protocol_path)}), "
            f"но пустой или нечитаемый."
        )

    # --- Собираем блоки ---
    blocks = build_prompt_blocks(
        protocol_text=protocol_text,
        session_name=session_name,
        session_date=session_date,
        template=template,
    )

    # --- Пишем файл ---
    path = write_prompt_file(
        session_dir=session_dir,
        blocks=blocks,
        fmt=fmt,
    )
    if not path:
        if not result["warning"]:
            result["warning"] = "Не удалось сохранить файл промпта"
        return result

    result["ok"] = True
    result["path"] = path
    return result